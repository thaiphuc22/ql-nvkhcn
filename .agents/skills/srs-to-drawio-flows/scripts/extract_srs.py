from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.document import Document as DocxDocument
    from docx.oxml.ns import qn
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph
except ImportError as exc:
    raise SystemExit("python-docx is required: pip install python-docx") from exc


FUNCTION_RE = re.compile(
    r"^\s*(?:function|chức\s*năng)\s*(\d+(?:\.\d+)*)\s*(?:[-–—:.)]\s*)?(.*)$",
    re.IGNORECASE,
)
EXCEPTION_RE = re.compile(
    r"\b(nếu|ngoại lệ|trường hợp|không có|không tồn tại|không hợp lệ|"
    r"trùng|xung đột|timeout|thất bại|bắt buộc|403|404|409|500|"
    r"đã thay đổi|vượt quá|dừng luồng)\b",
    re.IGNORECASE,
)
API_RE = re.compile(
    r"\b(GET|POST|PUT|PATCH|DELETE)\s+/|"
    r"\b[a-z][a-z0-9_]*(?:\.[a-z][a-z0-9_]*)?\b",
    re.IGNORECASE,
)
SECTION_NAMES = {
    "màn hình",
    "luồng nghiệp vụ",
    "mô tả chi tiết luồng nghiệp vụ",
    "mô tả chi tiết các thành phần",
    "thông tin chung chức năng",
}
COMMON_SECTION_RE = re.compile(
    r"thiết\s*kế.*dùng\s*chung|tái\s*sử\s*dụng|\bcommon\b",
    re.IGNORECASE,
)


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def iter_blocks(parent):
    if isinstance(parent, DocxDocument):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError(f"Unsupported parent: {type(parent)!r}")
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def paragraph_image_ids(paragraph: Paragraph) -> list[str]:
    result: list[str] = []
    for blip in paragraph._p.xpath(".//a:blip"):
        rid = blip.get(qn("r:embed"))
        if rid:
            result.append(rid)
    return result


def table_rows(table: Table) -> list[list[str]]:
    return [[clean(cell.text) for cell in row.cells] for row in table.rows]


def function_heading(text: str) -> tuple[str, str] | None:
    match = FUNCTION_RE.match(text)
    if not match:
        return None
    number, title = match.group(1), clean(match.group(2))
    return number, title or clean(text)


def has_function_info(records: list[dict], position: int) -> bool:
    for record in records[position + 1 : position + 25]:
        if record["type"] == "paragraph":
            text = clean(record["text"])
            if text.casefold() == "thông tin chung chức năng":
                return True
            if (
                record.get("style", "").casefold() == "heading 4"
                and text
                and text.casefold() not in SECTION_NAMES
            ):
                return False
        else:
            first_column = {
                clean(row[0]).casefold()
                for row in record["rows"]
                if row and clean(row[0])
            }
            if {"mục đích", "người dùng mục tiêu"} & first_column:
                return True
    return False


def function_candidates(
    records: list[dict], include_common: bool = False
) -> list[tuple[dict, str, str]]:
    output: list[tuple[dict, str, str]] = []
    current_heading2 = ""
    for position, record in enumerate(records):
        if record["type"] != "paragraph":
            continue
        text = clean(record["text"])
        if record.get("style", "").casefold() == "heading 2" and text:
            current_heading2 = text
        if not include_common and COMMON_SECTION_RE.search(current_heading2):
            continue
        explicit = function_heading(text)
        if explicit:
            output.append((record, explicit[0], explicit[1]))
            continue
        if (
            record.get("style", "").casefold() == "heading 4"
            and text
            and text.casefold() not in SECTION_NAMES
            and has_function_info(records, position)
        ):
            output.append((record, "", text))
    return output


def id_for(number: str, ordinal: int) -> str:
    return f"F{ordinal:02d}"


def context_heading(headings: list[dict], index: int, style: str) -> str:
    matches = [h for h in headings if h["index"] < index and h["style"].casefold() == style.casefold()]
    return matches[-1]["text"] if matches else ""


def first_info_table(chunk: list[dict]) -> dict[str, str]:
    known = {
        "mã chức năng",
        "mục đích",
        "người dùng mục tiêu",
        "đường dẫn truy cập",
        "tiền điều kiện",
        "hậu điều kiện",
    }
    for record in chunk:
        if record["type"] != "table":
            continue
        output: dict[str, str] = {}
        for row in record["rows"]:
            if len(row) < 2:
                continue
            key, value = clean(row[0]), clean(row[1])
            if key.casefold() in known:
                output[key.casefold()] = value
        if output:
            return output
    return {}


def section_lines(chunk: list[dict], wanted: set[str]) -> list[str]:
    active = False
    lines: list[str] = []
    for record in chunk:
        if record["type"] == "paragraph":
            text = clean(record["text"])
            if text.casefold() in SECTION_NAMES:
                active = text.casefold() in wanted
                continue
            if active and text:
                lines.append(text)
        elif active:
            for row in record["rows"]:
                value = " | ".join(x for x in map(clean, row) if x)
                if value:
                    lines.append(value)
    return lines


def collect_api_data(chunk: list[dict]) -> list[str]:
    output: list[str] = []
    for record in chunk:
        values: list[str] = []
        if record["type"] == "paragraph":
            values.append(clean(record["text"]))
        else:
            values.extend(" | ".join(x for x in map(clean, row) if x) for row in record["rows"])
        for value in values:
            if value and API_RE.search(value) and (
                "/" in value
                or any(token in value.casefold() for token in ("api", "table", "bảng", "db", "csdl", "service"))
            ):
                output.append(value)
    return list(dict.fromkeys(output))


def exception_candidates(lines: list[str]) -> list[str]:
    return list(dict.fromkeys(line for line in lines if EXCEPTION_RE.search(line)))


def main(source: Path, output_dir: Path, include_common: bool = False) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    media_dir = output_dir / "media"
    media_dir.mkdir(exist_ok=True)
    document = Document(source)
    records: list[dict] = []
    headings: list[dict] = []

    for index, block in enumerate(iter_blocks(document), 1):
        if isinstance(block, Paragraph):
            text = clean(block.text)
            style = block.style.name if block.style else ""
            images: list[str] = []
            for rid in paragraph_image_ids(block):
                part = block.part.related_parts.get(rid)
                if part is None:
                    continue
                source_name = str(part.partname).lstrip("/").replace("/", "_")
                target = media_dir / source_name
                if not target.exists():
                    target.write_bytes(part.blob)
                images.append(source_name)
            record = {
                "index": index,
                "type": "paragraph",
                "style": style,
                "text": text,
                "images": images,
            }
            records.append(record)
            if text and (style.casefold().startswith("heading") or function_heading(text)):
                headings.append(record)
        else:
            records.append({"index": index, "type": "table", "rows": table_rows(block)})

    candidates = function_candidates(records, include_common=include_common)
    if not candidates:
        raise SystemExit("No function headings found. Expected headings such as 'Function 1 - Title'.")

    functions: list[dict] = []
    for ordinal, (record, number, title) in enumerate(candidates, 1):
        start = record["index"]
        end = candidates[ordinal][0]["index"] if ordinal < len(candidates) else records[-1]["index"] + 1
        chunk = [item for item in records if start <= item["index"] < end]
        info = first_info_table(chunk)
        all_lines: list[str] = []
        images: list[str] = []
        for item in chunk:
            if item["type"] == "paragraph":
                if clean(item["text"]):
                    all_lines.append(clean(item["text"]))
                images.extend(item.get("images", []))
            else:
                all_lines.extend(
                    " | ".join(x for x in map(clean, row) if x)
                    for row in item["rows"]
                    if any(clean(x) for x in row)
                )
        flow_lines = section_lines(chunk, {"luồng nghiệp vụ", "mô tả chi tiết luồng nghiệp vụ"})
        ui_lines = section_lines(chunk, {"màn hình", "mô tả chi tiết các thành phần"})
        functions.append(
            {
                "id": id_for(number, ordinal),
                "number": number,
                "title": title,
                "module": context_heading(headings, start, "Heading 2"),
                "submodule": context_heading(headings, start, "Heading 3"),
                "source_start": start,
                "source_end": end,
                "code": info.get("mã chức năng", ""),
                "purpose": info.get("mục đích", ""),
                "actors": info.get("người dùng mục tiêu", ""),
                "path": info.get("đường dẫn truy cập", ""),
                "precondition": info.get("tiền điều kiện", ""),
                "postcondition": info.get("hậu điều kiện", ""),
                "images": list(dict.fromkeys(images)),
                "ui_lines": ui_lines,
                "flow_lines": flow_lines or all_lines,
                "api_data_candidates": collect_api_data(chunk),
                "exception_candidates": exception_candidates(flow_lines or all_lines),
            }
        )

    metadata = {
        "source": str(source),
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "inline_shapes": len(document.inline_shapes),
        "blocks": len(records),
        "function_count": len(functions),
    }
    (output_dir / "records.json").write_text(
        json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (output_dir / "functions.json").write_text(
        json.dumps(functions, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (output_dir / "meta.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(metadata, ensure_ascii=False))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output_dir", type=Path)
    parser.add_argument(
        "--include-common",
        action="store_true",
        help="Include functions defined under shared/common design sections.",
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    main(args.input, args.output_dir, include_common=args.include_common)
