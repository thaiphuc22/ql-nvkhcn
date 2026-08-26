#!/usr/bin/env python
"""Trích xuất nội dung 6 file Office trong `docs/hr_tool/` ra text đọc được.

Lý do tồn tại: các file gốc là .xlsx/.docx nhị phân — không grep được, không diff được,
không đọc được trong review. Các file .md cạnh script này là kết quả đã trích xuất bằng tay
từ output của script; chạy lại script khi khách gửi bản cập nhật để đối chiếu.

    cd docs/hr_tool/trich-xuat
    python trich-xuat.py            # in toàn bộ ra stdout
    python trich-xuat.py --sheets   # chỉ liệt kê tên sheet

⚠ BẮT BUỘC đặt PYTHONUTF8=1 trên Windows, nếu không print() sẽ chết với UnicodeEncodeError
  (cp1252 không encode được tiếng Việt có dấu):

    PYTHONUTF8=1 python trich-xuat.py

Phụ thuộc: openpyxl, python-docx (đã có sẵn trong Python của máy dev, không thêm vào repo).
"""

import sys
from pathlib import Path

import docx
import openpyxl
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

GOC = Path(__file__).resolve().parent.parent

XLSX = [
    "2025.08.01_HR_tool_Khach hang gui.xlsx",
    "Mo ta phan mem HR_23.06.2025.xlsx",
    "Book1.xlsx",
]
DOCX = [
    "VHT_Phan tich bai toan Phan Bo Nhan Cong.docx",
    "NS_Câu hỏi khảo sát nghiệp vụ hệ thống.docx",
    # "Phan Bo Nhan Cong Brd.docx" — trùng byte-for-byte với file trên, bỏ qua.
]


def in_docx(path: Path) -> None:
    """In cả đoạn văn lẫn bảng theo đúng thứ tự xuất hiện trong tài liệu.

    python-docx không cho duyệt hỗn hợp sẵn — `doc.paragraphs` và `doc.tables` là hai danh
    sách rời, ghép lại thì mất thứ tự. Phải tự duyệt cây XML của body.
    """
    d = docx.Document(str(path))
    for child in d.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = Paragraph(child, d).text.strip()
            if text:
                print(text)
        elif child.tag == qn("w:tbl"):
            for row in Table(child, d).rows:
                cells, da_thay = [], set()
                for c in row.cells:
                    # Ô merge ngang trả về cùng một _tc nhiều lần — lọc trùng, nếu không
                    # mỗi ô gộp sẽ bị in lặp đúng bằng số cột nó chiếm.
                    if id(c._tc) in da_thay:
                        continue
                    da_thay.add(id(c._tc))
                    cells.append(" ".join(c.text.split()))
                if any(cells):
                    print(" | ".join(cells))
            print("---")


def in_xlsx(path: Path, chi_ten_sheet: bool) -> None:
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    for ws in wb.worksheets:
        if chi_ten_sheet:
            print(f"  - {ws.title!r} {ws.max_row}x{ws.max_column} ({ws.sheet_state})")
            continue
        print(f"\n########## SHEET: {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            cells = ["" if c is None else str(c).replace("\n", " ⏎ ").strip() for c in row]
            while cells and cells[-1] == "":
                cells.pop()  # bỏ đuôi rỗng, nếu không mỗi dòng kéo dài tới cột cuối cùng
            if cells:
                print(f"{i}| " + " | ".join(cells))
    wb.close()


def main() -> int:
    chi_ten_sheet = "--sheets" in sys.argv
    for ten in XLSX:
        print(f"\n=== {ten}")
        in_xlsx(GOC / ten, chi_ten_sheet)
    if not chi_ten_sheet:
        for ten in DOCX:
            print(f"\n=== {ten}")
            in_docx(GOC / ten)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
